# HW1a 
# 1. Read dataset and interpolate missing values
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

# Read Excel file
file_path = 'HSI.xlsx'
data = pd.read_excel(file_path, index_col='Date', parse_dates=True)

# Interpolate missing entries using spline method of order 3
interpolated_data = data.interpolate(method='spline', order=3)
interpolated_data

# 2. Weekly and Monthly Dataframes
# Weekly dataframe
weekly_data = interpolated_data.resample('W-FRI').mean()

# Monthly dataframe
monthly_data = interpolated_data.resample('ME').mean()

print(weekly_data)
print(monthly_data)

# 3. Daily, Weekly, and Monthly returns
daily_returns = interpolated_data.pct_change()
weekly_returns = weekly_data.pct_change()
monthly_returns = monthly_data.pct_change()

print(daily_returns)
print(weekly_returns)
print(monthly_returns)

# 4. Covariance matrix for weekly data
cov_matrix = weekly_data.cov()
cov_matrix.to_csv('covHSI.csv')

print(cov_matrix)

# 5. Histogram for tencent (700 HK)
import matplotlib.pyplot as plt
from docx import Document

# Only use tencent values and drop the NA values
tencent_returns = daily_returns['700 HK'].dropna()

# Plot histogram
plt.hist(tencent_returns, bins=100, density=True)
plt.title('700 HK', fontsize=10)
plt.xlabel('Daily Returns')
plt.ylabel('Density')

# Save histogram as png file
plt.savefig('tencent_histogram.png')

# Create a Word document and insert the histogram
doc = Document()
doc.add_heading('Tencent Daily Returns Histogram', level=1)
doc.add_picture('tencent_histogram.png')
doc.save('Graph.doc')

# Show the plot
plt.show()

# 6. Standard deviation loop
# Empty list to store multiple dataframes
monthly_volatility_list = []

# Loop through each stock and each month
for stock in daily_returns.columns:
    for month in daily_returns.index.month.unique():
        for year in daily_returns.index.year.unique():
            month_data = daily_returns[(daily_returns.index.month == month) & (daily_returns.index.year == year)][stock]
            std_dev = month_data.std()
            # Temp dataframe
            temp_df = pd.DataFrame({
                'Stock': [stock],
                'Year': [year],
                'Month': [month],
                'StdDev': [std_dev]
            })
            # Append datafram to list
            monthly_volatility_list.append(temp_df)

# Concatenate all DataFrames in the list into a single DataFrame
monthly_volatility = pd.concat(monthly_volatility_list)

# Output the result to a CSV file
monthly_volatility.to_csv('HSI_vol.csv')

# 7. Question 6 but using 'resample' function
# Create empty dataframe this time 
monthly_volatility = pd.DataFrame()

for stock in daily_returns.columns:
    # Use resample function 
    monthly_stdev = daily_returns[stock].resample('ME').std()

    #Temp dataframe
    temp_df = monthly_stdev.reset_index()
    temp_df.columns = ['Date', 'StdDev']  
    temp_df['Stock'] = stock  
    temp_df['Year'] = temp_df['Date'].dt.year 
    temp_df['Month'] = temp_df['Date'].dt.month

    monthly_volatility = pd.concat([monthly_volatility, temp_df[['Stock', 'Year', 'Month', 'StdDev']]])

monthly_volatility.to_csv('HSI_vol.csv')

